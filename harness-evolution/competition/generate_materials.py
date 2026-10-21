#!/usr/bin/env python3
"""Create 16:9 slides and readable previews from the public evidence JSON."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BG = '#0c1017'
WHITE = '#edf2f7'
MUTED = '#93a4b6'
GREEN = '#a9f06e'
BLUE = '#71baff'
ORANGE = '#ffd28b'
FONT_FILE = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
FONT = 'Noto Sans CJK SC'


def box(x, y, w, h, text, size=34, color=WHITE):
    return dict(x=x, y=y, w=w, h=h, text=text, size=size, color=color)


def content(data):
    gains = '  ·  '.join(f"{r['task_id']} {r['verified_speedup']:.2f}×" for r in data['operators'])
    application = data['application']
    application_text = application.get('summary', 'GPU端到端结果尚未测量；不以单算子加速比推断整体收益。')
    campaign = data.get('campaign', {})
    series = data.get('measurement_series', {})
    if series.get('closed_arm_count') == 8:
        campaign_text = (f"160位置完整保留；原评分门未通过。\n"
                         f"DEV几何均值比{series['development_geomean_ratio']:.4f}，"
                         f"95%区间[{series['development_paired_ci_low']:.4f}, {series['development_paired_ci_high']:.4f}]。")
    else:
        campaign_text = campaign.get('summary', '干净版本对照待完成，当前不能声称搜索性能进化成立。')
    comparison = series.get('comparison', {'parent_source_commit': '54b05f2', 'candidate_source_commit': '7e9aaa1', 'gpu_indices': [0,3,4]})
    comparison_text = (f"原始{comparison['parent_source_commit'][:7]}  ↔  改进{comparison['candidate_source_commit'][:7]}\n"
        f"三组开发配对 + 独立留出 · 每轮完整20任务\n相同模型、种子、单任务预算和干净GPU{'/'.join(map(str, comparison['gpu_indices']))}")
    source_text = 'G001  编辑失败：先读取，整批编辑未应用\nG002  审批拒绝：调用未执行，仍需授权\nG003  上下文压缩：保留runtime_context类型'
    has_g004 = any(row['generation'] == 'G004' for row in data['source_changes'])
    if has_g004:
        source_text += '\nG004  路径拒绝：当前cwd与相对路径恢复提示'
    has_g005 = any(row['generation'] == 'G005' for row in data['source_changes'])
    if has_g005:
        source_text += '\nG005  拒绝后确认状态，作业失败后先诊断'
    has_g006 = any(row['generation'] == 'G006' for row in data['source_changes'])
    if has_g006:
        source_text += '\nG006  小窗口压缩预留按实际context约束'
    has_g007 = any(row['generation'] == 'G007' for row in data['source_changes'])
    if has_g007:
        source_text += '\nG007  结构化JSON模型视图压缩，保留原始事件'
    source_note = 'G003 CPU路径summary 2→1；' + ('G004 229项回归通过\n175文件新运行包，39次localhost原生请求资格通过' if has_g004 else '付费搜索收益待对照')
    if has_g005:
        source_note = 'G005 189项回归、551项guard通过\n175文件新运行包与39次localhost请求；搜索结果见原评分门'
    if has_g006:
        source_note = 'G006 221项相关回归、551项guard通过\n175文件独立运行包、39次localhost请求资格\n源码工程修复已验证；搜索收益须新版本对照'
    if has_g007:
        source_note = 'G007 364项相关回归、551项guard通过\n175文件独立新运行包、39次实际SDK请求\n搜索收益仍待V17；最近V16原晋级门未通过'
    has_g008 = any(row['generation'] == 'G008' for row in data['source_changes'])
    if has_g008:
        source_text += '\nG008  优先使用反馈中已有的异常诊断'
        source_note = 'G008 368项相关回归、551项guard通过，保留2项跳过\n175文件独立运行包、78次实际localhost请求\nV18完整对照未晋级；G008收益须接受新的完整对照'
    if any(row['generation'] == 'G010' for row in data['source_changes']):
        source_text = ('G001–G005  编辑、审批、工作区与作业恢复提示\n'
                       'G006  按实际context窗口预留压缩额度\n'
                       'G007–G008  JSON模型视图与已有异常诊断\n'
                       'G009  实际剩余原生请求/候选预算进入正文\n'
                       'G010  合法的结构化数值失败优先引导修复')
        source_note = ('源码作者：外层Codex；保留原权限/容差/24请求/四候选\n'
                       'G009与G010各39次实际本地SDK资格；175文件新运行包\n'
                       '工程资格通过，H0仍为incumbent，未证明RSI')
        ablation = data['closed_ablation']
        campaign_text = (f"V21 DEV {series['development_geomean_ratio']:.4f}，95%CI "
                         f"[{series['development_paired_ci_low']:.4f},{series['development_paired_ci_high']:.4f}]，未过门。\n"
                         f"V22 G009→G010正确8→12；全20比{ablation['development_geomean_ratio']:.4f}。\n"
                         f"GPU耗时比{ablation['original_promotion']['pairs'][0]['cost_ratios']['gpu_seconds']:.4f}>1.1；未晋级。")
        comparison_text = (f"H0 {comparison['parent_source_commit'][:7]} ↔ G009 {comparison['candidate_source_commit'][:7]}\n"
                           "三组DEV配对 + 独立留出；160位置保留\n"
                           "评测种子固定，模型采样未固定；G010另有单DEV消融")
    if any(row['generation'] == 'G012' for row in data['source_changes']):
        source_text = ('G001–G008  编辑、工作区、诊断与模型视图\n'
                       'G009  实际原生预算进入模型正文\n'
                       'G010  合法数值失败优先修复\n'
                       'G011  静态编译提示；V24未过门\n'
                       'G012  正确但未晋级的观测速度反馈')
        source_note = ('源码作者外层Codex；G012直接基于G010\n'
                       '437相关回归 / 175文件运行包 / 39次SDK资格\n'
                       'G012工程资格；搜索结论以当前闭账对照为准')
        ablation = data['closed_ablation']
        gpu_ratio = ablation['original_promotion']['pairs'][0]['cost_ratios']['gpu_seconds']
        cost_state = '成本门通过，正确性与CI门未过' if gpu_ratio <= 1.1 else '成本门未过'
        campaign_text = (f"V23 DEV {series['development_geomean_ratio']:.4f}，95%CI "
                         f"[{series['development_paired_ci_low']:.4f},{series['development_paired_ci_high']:.4f}]，未过门。\n"
                         f"V24 G010→G011正确11→10；K15缺评测，全20比{ablation['development_geomean_ratio']:.4f}。\n"
                         f"GPU耗时比{gpu_ratio:.4f}；{cost_state}。")
        comparison_text = (f"H0 {comparison['parent_source_commit'][:7]} ↔ G010 {comparison['candidate_source_commit'][:7]}\n"
                           'V23三组DEV + 独立留出；160位置保留\n'
                           'V24两组40位置闭账；V25 G012性能待证明')
    if any(row['generation'] == 'G013' for row in data['source_changes']):
        source_text = ('G001–G009  编辑、工作区、诊断、模型视图与预算\n'
                       'G010  合法数值失败优先修复\n'
                       'G011  静态编译提示；V24未过门\n'
                       'G012  未晋级候选观测反馈；V25未过门\n'
                       'G013  首次待完成观测留档，原超时内最多续查一次')
        source_note = ('源码作者外层Codex；G013直接基于G010\n'
                       '525相关回归 / 175文件运行包 / 48次本地SDK正文\n'
                       '合成请求3→2；V26性能待实测，RSI未证明')
        ablation = data['closed_ablation']
        pair = ablation['original_promotion']['pairs'][0]
        costs = pair['cost_ratios']
        parent = next(a for a in ablation['arms'] if a['arm'] == 'parent')
        candidate = next(a for a in ablation['arms'] if a['arm'] == 'candidate')
        revision = ablation['comparison']['campaign_revision']
        campaign_text = (f"V{revision}全20比{ablation['development_geomean_ratio']:.4f}，95%CI"
                         f"[{ablation['development_paired_ci_low']:.4f},{ablation['development_paired_ci_high']:.4f}]\n"
                         f"token/USD/GPU成本比{costs['token_upper_bound']:.4f}/{costs['usd_upper_bound']:.4f}/{costs['gpu_seconds']:.4f}\n"
                         '原正确性、置信区间与成本门未过；H0保持，RSI未证明')
        comparison_text = ('V23 H0/G010：160位置、三组DEV与独立留出\n'
                           f"V{revision} {ablation['comparison']['parent_generation']}→{ablation['comparison']['candidate_generation']}：40位置，正确{parent['correct_count']}→{candidate['correct_count']}\n"
                           'V26 G010/G013：原生20路两组；性能待实测')
        if revision == 26:
            source_note = ('源码作者外层Codex；G013直接基于G010\n'
                           '525相关回归 / 175文件运行包 / 48次本地SDK正文\n'
                           'V26原对照闭账，无晋级；RSI未证明')
            campaign_text = (f"V26全20比{ablation['development_geomean_ratio']:.4f}，95%CI"
                             f"[{ablation['development_paired_ci_low']:.4f},{ablation['development_paired_ci_high']:.4f}]\n"
                             f"token/USD/GPU成本比{costs['token_upper_bound']:.4f}/{costs['usd_upper_bound']:.4f}/{costs['gpu_seconds']:.4f}\n"
                             '完整评测、置信区间与GPU成本门未过；H0保持，RSI未证明')
            comparison_text = ('V23 H0/G010：160位置、三组DEV与独立留出\n'
                               f"V26 G010→G013：40位置，正确{parent['correct_count']}→{candidate['correct_count']}\n"
                               f"G013最终评测{candidate['evaluation_complete_count']}/20；缺失K16保留，无晋级")
    if any(row['generation'] == 'G014' for row in data['source_changes']):
        source_text += '\nG014  仅对运行中作业续查；排队首个结果交还模型'
        source_note = ('源码作者外层Codex；G014基于G013\n'
                       '526相关回归 / 175运行文件 / 52次SDK正文\n'
                       '排队路径请求3对2；付费待测，RSI未证明')
    if any(row['generation'] == 'G017' for row in data['source_changes']):
        source_text = ('G001–G009  工具恢复、模型视图与实际预算\n'
                       'G010–G012  数值、编译及观测反馈；失败保留\n'
                       'G013–G014  原超时内有界续查，保留排队取舍\n'
                       'G016  形状回退提示；V29未过原准入\n'
                       'G017  重复实现的可信回执反馈，保留实际写入')
        source_note = ('源码作者外层Codex；G017从exact G014独立分支\n'
                       '546项相关检查 / 175运行文件 / 67次SDK正文\n'
                       'V30待四卡GPU准入与付费实测；RSI未证明')
        ablation = data['closed_ablation']
        costs = ablation['original_promotion']['pairs'][0]['cost_ratios']
        parent = next(a for a in ablation['arms'] if a['arm'] == 'parent')
        candidate = next(a for a in ablation['arms'] if a['arm'] == 'candidate')
        campaign_text = (f"V29全20比{ablation['development_geomean_ratio']:.4f}，95%CI"
                         f"[{ablation['development_paired_ci_low']:.4f},{ablation['development_paired_ci_high']:.4f}]\n"
                         f"token/USD/GPU成本比{costs['token_upper_bound']:.4f}/{costs['usd_upper_bound']:.4f}/{costs['gpu_seconds']:.4f}\n"
                         '完整评测、置信区间与GPU成本门未过；H0保持，RSI未证明')
        comparison_text = ('V23 H0/G010：160位置、三组DEV与独立留出\n'
                           f"V29 G014→G016：四卡40位置，正确{parent['correct_count']}→{candidate['correct_count']}\n"
                           f"G016最终评测{candidate['evaluation_complete_count']}/20；缺失K16保留，未晋级")
    cached = data.get('cached_application', {})
    packed = data.get('packed_application', {})
    def range_text(rows, key, digits=3):
        values = [row[key]['median'] for row in rows]
        return f'{min(values):.{digits}f}–{max(values):.{digits}f}×'
    sequence_text = application_text
    if application.get('cases'):
        sequence_text = '序列：相对eager ' + range_text(application['cases'], 'against_eager', 2) + '；相对编译库 ' + range_text(application['cases'], 'against_compiled', 2)
    cached_text = 'KV缓存：强库对照 ' + range_text(cached['cases'], 'against_strong_library') if cached.get('cases') else 'KV缓存GPU结果待测'
    if data.get('independent_cached'):
        cached_text += '；GPU1→3独立复测完成'
    packed_text = ('合并投影：普通FP16库 ' + range_text(packed['cases'], 'against_fp16_library', 2) + '\n六库最强对照 ' + range_text(packed['cases'], 'against_strong_library') + '，未全形状胜出') if packed.get('cases') else '合并投影GPU结果待测'
    variant = data.get('packed_variant_application', {})
    if variant.get('replicas'):
        cases = [case for replica in variant['replicas'] for case in replica['cases']]
        packed_text = ('新GEMM / GPU0与3：普通FP16库 ' + range_text(cases, 'against_fp16_library', 2)
                       + '\n六库最强 ' + range_text(cases, 'against_strong_library') + '，边界更慢')
    pretrained_text = ('真实GPT-2：短序列约2.02×；长序列数值失败保留' if data.get('pretrained_application')
                       else '相同固定权重与输入 · 30轮配对 · 不代表预训练文本质量')
    if data.get('pretrained_precision_application'):
        pretrained_text = 'GPT-2精度诊断：仅短形状通过；强库约1.000×，长形状仍失败'
    if data.get('pretrained_numerical_diagnosis'):
        pretrained_text = 'GPT-2实际GPU定位：局部48/48通过；全模型1/4通过，误差逐层累积'
    if data.get('welford_application'):
        pretrained_text = 'Welford：四形状数值通过，普通库0.90–0.95×；新编译仅1/4通过，未计时'
    slides = [
        ('KernelEvolve', [box(115,250,1660,240,'让GPU算子优化与AGH改进\n都有可复验证据',76),
            box(120,545,1630,120,'20路AGH搜索  /  轨迹诊断  /  源码修改  /  独立复测',34,MUTED),
            box(120,780,1650,90,'AGNES HACKATHON · 硕博组 · 模型 mimo-v2.6-pro',27,GREEN)],
         '0:00–0:25 介绍项目：把GPU算子搜索和智能体运行改进连接成可审查的实验闭环。模型是mimo-v2.6-pro，底座为AGH，源码改动由外层开发者完成。'),
        ('真正的问题：一次快，不代表系统变强', [box(120,240,1660,170,'并发搜索会产生代码、错误、成本和状态。\n没有可信反馈，优化结果就难以复现。',48),
            box(120,500,1650,245,'我们把三个问题分开验证：\n① 生成的算子是否正确且更快\n② 修改AGH之后，搜索是否稳定改善\n③ 完整应用是否获得实际收益',38,MUTED)],
         '0:25–0:50 科研与工程的GPU优化需要正确性、性能和可重复执行。搜索一次找出冠军不等于智能体进化，单算子变快也不等于应用变快。三个问题采用独立证据。'),
        ('AGH完成原生工具与反馈闭环', [box(120,250,1650,200,'获取任务 → 写入/编辑源码 → 提交候选\n→ GPU评测 → 读取反馈 → 修订源码',48,GREEN),
            box(120,535,1650,170,'20任务并发 · 每任务4候选 · 24次原生模型请求\n5个模型槽位统一MiMo/off · 10个worker显式固定\n整批严格关闭后，外层才修改AGH并冻结下一代',35,MUTED)],
         '0:50–1:20 展示真实AGH执行链：任务、代码、GPU反馈和修改。原生trace绑定候选源文件与队列job，不把终端退出等同于整批关闭。全批清理后开始下一代源码修改。'),
        (f"已有独立复验：{len(data['operators'])}个算子", [box(120,240,1650,190,gains,42,GREEN),
            box(120,520,1650,170,'五类形状 · 固定项目基线 · 原文件SHA256\nCSR原大比值对Python基线；新增强库复测约3.358×\n弱收益形状和不同基线范围保留，不宣称全部强库胜出' if data.get('graph_application') else '五类形状 · 固定项目基线 · 原文件SHA256\nCSR约724×针对Python分段基线，跨设备尚未新增复现\n弱收益形状和失败均保留，不宣称胜过所有强库',36,MUTED)],
         '1:20–1:50 已展示的算子有原独立GPU验证。新增CSR约724倍只针对已资格验证的Python分段基线，尚未新增跨设备复验。展示形状几何均值、弱形状与源文件，不声称优于所有强库。'),
        ('AGH源码确实发生了变化', [box(120,225,1640,400,source_text,34),
            box(120,650,1650,215,source_note,31,ORANGE)],
         '1:50–2:20 每次改动都有具体源commit和工程回归。压缩消息类型修复减少一次CPU summary请求；路径恢复提示提供当前工作区和可操作路径。工程证据仍须接受付费搜索对照。'),
        ('同条件版本对照，验收搜索进化', [box(120,240,1650,210,comparison_text,40,GREEN),
            box(120,605,1650,215 if data.get('closed_ablation') else 160,campaign_text,31 if data.get('closed_ablation') else 36,ORANGE)],
         '2:20–2:50 预声明三组开发配对和独立留出，交错执行次序，杜绝继承冠军。全部20任务纳入分母，沿用原正确性、资源、置信区间和提升门。当前状态严格按实测显示。'),
        ('从算子到完整decoder计算块', [box(120,235,1650,145,'Q/K/V投影 → RoPE → 因果注意力\n→ 输出投影 → 门控MLP → 残差',40,GREEN),
            box(120,410,1650,135,sequence_text,34,ORANGE),
            box(120,560,1650,130,cached_text,32,ORANGE),
            box(120,705,1650,120,packed_text,32,ORANGE),
            box(120,850,1650,60,pretrained_text,27,MUTED)],
         '2:50–3:20 序列与KV缓存、合并投影保留原实测和强基线限制。Welford版本保留全部124M预训练参数，四形状完整输出逐位一致，但实际计时约为普通库的0.90至0.95倍，没有端到端加速。另一个共同编译配置消除了图中断，仅一形状数值通过，未计时。'),
        ('严谨性：失败、未知与污染都有记录', [box(120,235,1650,190,f"V21全160位置 · 旧阶段八卡206个真实job claims\nV22闭账未知预留{data['closed_history_v22']['unknown_all_history']}，不等于确认账单" if data.get('closed_history_v22') else f"本轮160位置 · 旧阶段八卡206个真实job claims\n未知请求{data['accounting']['unknown_requests']}，保守费用上界并非确认账单",41),
            box(120,540,1650,190,'驻留外部进程不会被清理，原污染判据保留\n历史G003早停根因未解析，新监控记录精确字段码\n源码、运行包、队列、原始计时与证据按SHA256绑定',36,MUTED)],
         '3:20–3:45 原始失败和未知不删除，保守预留不退款。八卡参与执行已经有证据，但执行授权和有效测量分开判断。公开演示只展示安全字段，失败事件保持可见。'),
        ('交付可以独立核查的优化系统', [box(120,240,1660,190,'原生AGH调用链 + 版本化源码\n可复现算子与完整应用 + 离线证据演示',48,GREEN),
            box(120,560,1650,170,'打开每个结果的源文件、形状表现与失败记录。\n重复对照通过才保留性能进化结论；\n改动无效就回到整批轨迹，继续修改、验证和实测。',36,MUTED)],
         '3:45–4:00 结束展示。贡献是可验证的搜索、源码改进与应用评测系统。结论以实际数据为准，不把开发者辅助修改称为AGH原生自改，不预设获奖或性能成立。'),
    ]
    if data.get('graph_application'):
        graph = data['graph_application']
        gains = '  /  '.join(f"GPU{r['gpu']} {r['strong_gain']['median']:.4f}×" for r in graph['replicas'])
        slides[6] = ('真实引文图计算，强库对照约1.59×', [
            box(120,220,1650,130,'ogbn-arxiv：169343节点 · 1166243边 · 两层FP32图计算',42,GREEN),
            box(120,375,1650,150,gains+'\n相对每轮五种合格库的最快实现，耗时下降约37%',40,GREEN),
            box(120,555,1650,155,'全量 / 2048节点窗口 / 边界，全部通过原数值门\n1080 GPU + 1080主机原样本；30轮配对与95%CI保留',33,MUTED),
            box(120,740,1650,100,'源码作者Codex；固定随机权重未训练，不证明分类精度或AGH进化',30,ORANGE),
            box(120,875,1650,70,'GPT-2失败也保留：Welford普通库0.90–0.95×；新编译仅1/4数值通过',27,MUTED)],
            '2:50–3:20 图计算采用真实引文图数据、相同输入与FP32权重。两设备全量强库比约1.59倍，窗口和边界也通过原数值门。所有合格后端、逐轮计时、置信区间与源码可核查。权重未训练；应用由外层Codex实现，不能据此证明AGH自改或搜索进化。旧GPT-2的数值失败和性能损失继续保留。')
        # Current CSR reproduction is a distinct scope from its original report.
        slides[3] = (slides[3][0], slides[3][1], '1:20–1:50 展示原算子文件、五种形状和独立验证。CSR原大比值针对Python分段基线；新增复测一张卡对合格编译库约3.358倍，另一张卡仅对Python基线约730.74倍，基线范围不同。不能相乘单算子比值推断整模型收益。')
    if data.get('closed_history_v24'):
        history = data['closed_history_v24']
        slides[7][1][0]['text'] = (f"V23全160位置 / V24全40位置保留\n"
                                  f"V24历史{history['request_rows']}请求，未知预留{history['unknown_all_history']}；不是确认账单")
    if data.get('closed_history_v25'):
        history = data['closed_history_v25']
        slides[7][1][0]['text'] = (f"V23全160位置 / V24与V25各40位置保留\n"
                                  f"V25历史{history['request_rows']}请求，未知预留{history['unknown_all_history']}；不是确认账单")
        slides[5] = (slides[5][0], slides[5][1], slides[5][2] +
            ' V24的G011正确11→10、K15缺评测保留；V25的G012正确13→10和成本失败单独登记。G012的437项回归与39次SDK资格没有证明搜索晋级；G013的48次正文验证同样只证明执行行为。')
    if data.get('closed_history_v26'):
        history = data['closed_history_v26']
        slides[7][1][0]['text'] = (f"V23全160位置 / V24至V26各40位置保留\n"
                                  f"V26历史{history['request_rows']}请求，未知预留{history['unknown_all_history']}；不是确认账单")
        slides[5] = (slides[5][0], slides[5][1], slides[5][2] +
            ' V26两版正确11/20，G013最终评测19/20、缺失K16保留；置信区间含1，GPU成本比1.1276超过原1.1门，无晋级。运行中20次续查10次拿到结果，排队20次仅1次；观测数不等于节省的模型请求。')
        if any(row['generation'] == 'G014' for row in data['source_changes']):
            slides[4] = (slides[4][0], slides[4][1], slides[4][2] +
                ' G014只对运行中的作业自动续查，排队结果交还模型。真实SDK控制中运行路径两版均2次请求，排队路径G014为3、G013为2；这项取舍必须经过付费对照，尚无收益或晋级证据。')
    if data.get('closed_history_v29'):
        history = data['closed_history_v29']
        slides[7][1][0]['text'] = (f"V23全160位置 / V29四卡40位置保留\n"
                                  f"V29历史{history['request_rows']}请求，未知预留{history['unknown_all_history']}；不是确认账单")
        slides[4] = (slides[4][0], slides[4][1],
            '1:50–2:20 G017从exact G014独立分支，不继承未接受G016。重复实现提示确实进入下一模型请求；'
            '原写入和hook保留，父子复制流程均5请求、4次write。67个本地CPU HTTP体只证明机制，V30四卡付费未实测。')
        slides[5] = (slides[5][0], slides[5][1],
            '2:20–2:50 V29两组40位置全部保留：正确10到12，最终20到19、K16缺失，综合比0.978851。'
            '置信区间含1，GPU成本比1.444701，G016未接受。H0不变；只有后续重复、留出、新任务与两代晋级证据才能证明持续改进。')
    return slides


def wrapped_lines(text, font, width):
    lines = []
    for paragraph in text.split('\n'):
        line = ''
        for char in paragraph:
            if line and font.getlength(line + char) > width:
                lines.append(line)
                line = char
            else:
                line += char
        lines.append(line)
    return lines


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    data = json.loads(args.data.read_text())
    assert data['mode'] == 'recorded'
    if not Path(FONT_FILE).is_file():
        raise SystemExit('Noto Sans CJK font unavailable')
    args.out.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333333), Inches(7.5)
    notes = []
    for index, (title, elements, note) in enumerate(content(data), 1):
        elements = [box(120,60,1650,60,'KERNELEVOLVE  /  AGH × GPU  /  硕博组',24,GREEN),
                    box(120,133,1650,88,title,52), *elements,
                    box(120,995,1680,44,f'{index:02d} / 09    记录证据 · 非实时执行    mimo-v2.6-pro',23,MUTED)]
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(BG[1:])
        image = Image.new('RGB', (1920,1080), BG)
        draw = ImageDraw.Draw(image)
        draw.line((120,955,1790,955), fill='#293544', width=2)
        for element in elements:
            font = ImageFont.truetype(FONT_FILE, element['size'])
            lines = wrapped_lines(element['text'], font, element['w'])
            line_height = round(element['size'] * 1.45)
            if len(lines) * line_height > element['h']:
                raise ValueError(f"slide {index} text exceeds its layout: {element['text'][:35]}")
            draw.multiline_text((element['x'],element['y']), '\n'.join(lines), font=font,
                                fill=element['color'], spacing=round(element['size']*.45))
            shape = slide.shapes.add_textbox(Inches(element['x']/144), Inches(element['y']/144),
                                            Inches(element['w']/144), Inches(element['h']/144))
            frame = shape.text_frame
            frame.margin_left = frame.margin_top = frame.margin_right = frame.margin_bottom = Inches(0)
            frame.word_wrap = False
            for line_index, line in enumerate(lines):
                paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.45
                run = paragraph.add_run()
                run.text = line
                run.font.name = FONT
                run.font.size = Pt(element['size']/2)
                run.font.color.rgb = RGBColor.from_string(element['color'][1:])
                ea = OxmlElement('a:ea'); ea.set('typeface', FONT); run._r.get_or_add_rPr().append(ea)
        slide.notes_slide.notes_text_frame.text = note
        image.save(args.out / f'slide-{index:02d}.png')
        notes.append(f'{index}. {title}\n\n{note}\n')
    presentation.core_properties.title = 'KernelEvolve：AGH GPU优化与持续改进'
    presentation.core_properties.subject = 'AGNES Hackathon masters/PhD group'
    presentation.core_properties.comments = 'Source-bound recorded evidence; no invented RSI or application gains.'
    target = args.out / 'KernelEvolve-硕博组.pptx'
    presentation.save(target)
    (args.out / 'four-minute-script.txt').write_text('\n'.join(notes))
    proof = {'schema_version':1, 'data_sha256':hashlib.sha256(args.data.read_bytes()).hexdigest(),
             'slide_count':len(presentation.slides), 'aspect_ratio':'16:9',
             'pptx_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
             'previews':[f'slide-{i:02d}.png' for i in range(1,len(presentation.slides)+1)],
             'all_text_fits_preview':True, 'recorded':True,
             'rsi_performance_established':data['rsi_performance_established'],
             'application_status':data['application']['status']}
    (args.out / 'materials-manifest.json').write_text(json.dumps(proof,indent=2)+'\n')
    print(json.dumps(proof))


if __name__ == '__main__':
    main()
